"""DOCX parsing and HTML fragment generation."""

from __future__ import annotations

import html
import io
import logging
from typing import TYPE_CHECKING

from docx import Document
from docx.oxml.ns import qn as _docx_qn

from epub_converter.constants import (
    DEFAULT_IMAGE_ALT,
    _A_BLIP,
    _FOOTNOTES_RT,
    _R_EMBED,
    _R_HYP,
    _W_B,
    _W_FOOTNOTE_REF,
    _W_FN_ID,
    _W_HYPERLINK,
    _W_I,
    _W_P,
    _W_R,
    _W_RPR,
    _W_STRIKE,
    _W_T,
    _W_TBL,
    _W_U,
    _WP_DOC_PR,
)
from epub_converter.styles import (
    is_bullet_list,
    is_heading,
    is_numbered_list,
    is_quote_style,
)
from epub_converter.utils import xml_escape

if TYPE_CHECKING:
    from docx.text.paragraph import Paragraph

logger = logging.getLogger(__name__)


def _default_alt(lang: str) -> str:
    return DEFAULT_IMAGE_ALT.get(lang, DEFAULT_IMAGE_ALT["en"])


def _image_alt_from_run(r_elem) -> str:
    for doc_pr in r_elem.iter(_WP_DOC_PR):
        for attr in ("descr", "title"):
            val = doc_pr.get(attr)
            if val and val.strip():
                return val.strip()
    return ""


def _run_elem_to_html(
    r_elem,
    images: dict[str, bytes],
    part,
    default_alt: str,
) -> str:
    blips = r_elem.findall(f".//{_A_BLIP}")
    if blips:
        r_embed = blips[0].get(_R_EMBED)
        if r_embed and r_embed in part.rels:
            try:
                img_part = part.rels[r_embed].target_part
                blob = img_part.blob
                ct = img_part.content_type
                ext = ".png" if "png" in ct else ".jpg"
                img_id = f"img{len(images) + 1:03d}{ext}"
                images[img_id] = blob
                alt = _image_alt_from_run(r_elem) or default_alt
                return f'<img src="../images/{img_id}" alt="{xml_escape(alt)}"/>'
            except Exception as exc:
                logger.warning("Inline image extraction failed: %s", exc)
        return ""

    t_elems = r_elem.findall(_W_T)
    text = xml_escape("".join((t.text or "") for t in t_elems))
    if not text:
        return ""

    rPr = r_elem.find(_W_RPR)
    bold = rPr is not None and rPr.find(_W_B) is not None
    italic = rPr is not None and rPr.find(_W_I) is not None
    underline = rPr is not None and rPr.find(_W_U) is not None
    strike = rPr is not None and rPr.find(_W_STRIKE) is not None

    if bold and italic:
        text = f"<strong><em>{text}</em></strong>"
    elif bold:
        text = f"<strong>{text}</strong>"
    elif italic:
        text = f"<em>{text}</em>"
    if underline:
        text = f"<u>{text}</u>"
    if strike:
        text = f"<s>{text}</s>"
    return text


def para_to_html(
    para: Paragraph,
    images: dict[str, bytes],
    default_alt: str,
    fn_refs: list | None = None,
    footnotes: dict[str, str] | None = None,
) -> str:
    parts = []
    part = para.part
    for child in para._p:
        if child.tag == _W_R:
            ref_elem = child.find(f".//{_W_FOOTNOTE_REF}")
            if ref_elem is not None:
                fn_id = ref_elem.get(_W_FN_ID, "")
                if fn_refs is not None and footnotes is not None and fn_id in footnotes:
                    fn_refs.append((fn_id, footnotes[fn_id]))
                parts.append(f'<sup><a epub:type="noteref" href="#fn-{fn_id}">[{fn_id}]</a></sup>')
            else:
                parts.append(_run_elem_to_html(child, images, part, default_alt))
        elif child.tag == _W_HYPERLINK:
            r_id = child.get(_R_HYP)
            href = ""
            if r_id and r_id in part.rels:
                try:
                    href = xml_escape(part.rels[r_id].target_ref or "")
                except Exception as exc:
                    logger.debug("Hyperlink target read failed: %s", exc)
            inner = "".join(_run_elem_to_html(r, images, part, default_alt) for r in child.findall(_W_R))
            if inner:
                parts.append(f'<a href="{href}">{inner}</a>' if href else inner)
    return "".join(parts)


def table_to_html(table) -> str:
    rows = ["<table>"]
    for i, row in enumerate(table.rows):
        rows.append("<tr>")
        for cell in row.cells:
            tag = "th" if i == 0 else "td"
            rows.append(f"  <{tag}>{xml_escape(cell.text.strip())}</{tag}>")
        rows.append("</tr>")
    rows.append("</table>")
    return "\n".join(rows)


def extract_footnotes(doc) -> dict[str, str]:
    try:
        fn_part = doc.part.part_related_by(_FOOTNOTES_RT)
        result = {}
        for fn in fn_part._element.findall(_docx_qn("w:footnote")):
            fn_id = fn.get(_W_FN_ID, "")
            if fn_id in ("-1", "0"):
                continue
            texts = [xml_escape(t.text) for t in fn.findall(f".//{_W_T}") if t.text]
            result[fn_id] = "".join(texts)
        return result
    except Exception as exc:
        logger.debug("Footnotes not available: %s", exc)
        return {}


def parse_docx(docx_bytes: bytes, lang: str = "pl") -> tuple[list, dict[str, bytes]]:
    """Parse DOCX into chapters and inline images.

    chapters: list of (id, title, epub_type, role, body_html)
    """
    doc = Document(io.BytesIO(docx_bytes))
    footnotes = extract_footnotes(doc)
    default_alt = _default_alt(lang)
    chapters: list = []
    images: dict[str, bytes] = {}
    cur_id: str | None = None
    cur_title: str | None = None
    cur_body: list[str] = []
    cur_fn_refs: list = []
    in_ul = False
    in_ol = False
    chapter_counter = 0

    front_kw = {"wstep", "wstęp", "przedmowa", "wprowadzenie", "preface", "introduction", "foreword", "prolog"}
    back_kw = {"zakonczenie", "zakończenie", "epilog", "podsumowanie", "conclusion", "afterword", "epilogue"}

    para_map = {id(p._p): p for p in doc.paragraphs}
    table_map = {id(t._tbl): t for t in doc.tables}

    def close_list(parts: list[str]) -> None:
        nonlocal in_ul, in_ol
        if in_ul:
            parts.append("</ul>")
            in_ul = False
        if in_ol:
            parts.append("</ol>")
            in_ol = False

    def flush_chapter() -> None:
        if cur_title is None:
            return
        close_list(cur_body)
        if cur_fn_refs:
            seen: set = set()
            items: list[str] = []
            for fn_id, fn_text in cur_fn_refs:
                if fn_id not in seen:
                    seen.add(fn_id)
                    items.append(
                        f'<aside id="fn-{fn_id}" epub:type="footnote" role="note">'
                        f"<p><sup>[{fn_id}]</sup> {fn_text}</p></aside>"
                    )
            cur_body.append(
                '<section class="footnotes-section">\n' + "\n".join(items) + "\n</section>"
            )
        key = html.unescape(cur_title).lower().strip()
        if key in front_kw:
            epub_type, role = "frontmatter", "doc-preface"
        elif key in back_kw:
            epub_type, role = "backmatter", "doc-conclusion"
        else:
            epub_type, role = "chapter", "doc-chapter"
        chapters.append((cur_id, cur_title, epub_type, role, "\n".join(cur_body)))

    for child in doc.element.body:
        if child.tag == _W_P:
            para = para_map.get(id(child))
            if para is None:
                continue

            style_name = para.style.name if para.style else "Normal"
            text = para.text.strip()

            is_h1 = is_heading(style_name, 1)
            is_h2 = is_heading(style_name, 2)
            is_h3 = is_heading(style_name, 3)
            is_quote = is_quote_style(style_name)
            is_bullet = is_bullet_list(style_name)
            is_number = is_numbered_list(style_name)

            inner = para_to_html(
                para, images, default_alt, cur_fn_refs if cur_title else None, footnotes
            )
            fallback = xml_escape(text)

            if is_h1:
                flush_chapter()
                chapter_counter += 1
                cur_id = "ch" + str(chapter_counter).zfill(2)
                cur_title = xml_escape(text) if text else "Rozdzial"
                cur_body = []
                cur_fn_refs = []
                in_ul = in_ol = False
            elif cur_title is None:
                continue
            elif is_h2:
                close_list(cur_body)
                cur_body.append("<h2>" + (inner or fallback) + "</h2>")
            elif is_h3:
                close_list(cur_body)
                cur_body.append("<h3>" + (inner or fallback) + "</h3>")
            elif is_quote:
                close_list(cur_body)
                cur_body.append("<blockquote><p>" + (inner or fallback) + "</p></blockquote>")
            elif is_bullet:
                if not in_ul:
                    if in_ol:
                        cur_body.append("</ol>")
                        in_ol = False
                    cur_body.append("<ul>")
                    in_ul = True
                cur_body.append("<li>" + (inner or fallback) + "</li>")
            elif is_number:
                if not in_ol:
                    if in_ul:
                        cur_body.append("</ul>")
                        in_ul = False
                    cur_body.append("<ol>")
                    in_ol = True
                cur_body.append("<li>" + (inner or fallback) + "</li>")
            elif text or inner:
                close_list(cur_body)
                cur_body.append("<p>" + (inner or fallback) + "</p>")
            else:
                close_list(cur_body)

        elif child.tag == _W_TBL and cur_title is not None:
            table = table_map.get(id(child))
            if table is not None:
                close_list(cur_body)
                cur_body.append(table_to_html(table))

    flush_chapter()
    return chapters, images
