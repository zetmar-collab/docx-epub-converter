"""Helpers for building in-memory DOCX fixtures.

python-docx has no public API for hyperlinks or footnotes, so those are built
from raw OOXML here rather than being faked in the tests themselves - the goal
is to exercise the same element shapes Word actually produces.
"""

from __future__ import annotations

import io

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from PIL import Image

_W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_FOOTNOTES_CT = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
)


def new_document():
    """A blank document from python-docx's default template."""
    return Document()


def to_bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def add_run(para, text: str, **fmt) -> None:
    """Add a run with optional bold/italic/underline/strike formatting."""
    run = para.add_run(text)
    for attr in ("bold", "italic", "underline"):
        if fmt.get(attr):
            setattr(run.font, attr, True)
    if fmt.get("strike"):
        run.font.strike = True


def add_hyperlink(para, url: str, text: str) -> None:
    """Append a real w:hyperlink element backed by an external relationship."""
    r_id = para.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    para._p.append(link)


def add_footnote_reference(para, footnote_id: str) -> None:
    """Append a run containing only a w:footnoteReference, as Word does."""
    run = OxmlElement("w:r")
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), footnote_id)
    run.append(ref)
    para._p.append(run)


def attach_footnotes(doc, notes: dict[str, str]) -> None:
    """Attach a /word/footnotes.xml part holding the given id -> text notes.

    Word always emits the separator notes with ids -1 and 0; they are included
    so the parser's filtering of them is genuinely exercised.
    """
    body = ['<w:footnote w:id="-1"/>', '<w:footnote w:id="0"/>']
    for fn_id, text in notes.items():
        escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        body.append(
            f'<w:footnote w:id="{fn_id}"><w:p><w:r><w:t>{escaped}</w:t></w:r></w:p></w:footnote>'
        )
    xml = f"<w:footnotes {_W_NS}>{''.join(body)}</w:footnotes>"
    part = XmlPart(
        PackURI("/word/footnotes.xml"), _FOOTNOTES_CT, parse_xml(xml), doc.part.package
    )
    doc.part.relate_to(part, RT.FOOTNOTES)


def png_bytes(size=(12, 12), color=(200, 30, 30)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


def add_image(para, alt: str | None = None, descr_attr: str = "descr") -> None:
    """Add an inline image, optionally carrying alt text in wp:docPr."""
    run = para.add_run()
    run.add_picture(io.BytesIO(png_bytes()))
    if alt is not None:
        for doc_pr in run._r.iter(
            "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
        ):
            doc_pr.set(descr_attr, alt)


def simple_doc(paragraphs: list[tuple[str, str]]) -> bytes:
    """Build a document from (style_name, text) pairs."""
    doc = new_document()
    for style, text in paragraphs:
        doc.add_paragraph(text, style=style)
    return to_bytes(doc)
