"""Tests for DOCX -> HTML fragment conversion."""

import io

import pytest
from docx import Document

import docx_factory as F
from epub_converter.docx_parser import (
    extract_footnotes,
    parse_docx,
    table_to_html,
)


# ---------------------------------------------------------------------------
# Chapter splitting
# ---------------------------------------------------------------------------


def test_each_h1_starts_a_new_chapter():
    chapters, _ = parse_docx(
        F.simple_doc(
            [
                ("Heading 1", "Pierwszy"),
                ("Normal", "Tresc jeden"),
                ("Heading 1", "Drugi"),
                ("Normal", "Tresc dwa"),
            ]
        )
    )
    assert [c[1] for c in chapters] == ["Pierwszy", "Drugi"]
    assert "Tresc jeden" in chapters[0][4]
    assert "Tresc jeden" not in chapters[1][4]


def test_chapter_ids_are_zero_padded_and_sequential():
    chapters, _ = parse_docx(
        F.simple_doc([("Heading 1", f"R{i}") for i in range(1, 4)])
    )
    assert [c[0] for c in chapters] == ["ch01", "ch02", "ch03"]


def test_text_before_first_h1_is_dropped():
    chapters, _ = parse_docx(
        F.simple_doc(
            [
                ("Normal", "Sierocy akapit przed rozdzialem"),
                ("Heading 1", "Rozdzial"),
                ("Normal", "Wlasciwa tresc"),
            ]
        )
    )
    assert len(chapters) == 1
    assert "Sierocy" not in chapters[0][4]
    assert "Wlasciwa tresc" in chapters[0][4]


def test_document_without_h1_yields_no_chapters():
    chapters, _ = parse_docx(F.simple_doc([("Normal", "Sam tekst, zero naglowkow")]))
    assert chapters == []


def test_empty_h1_falls_back_to_placeholder_title():
    chapters, _ = parse_docx(
        F.simple_doc([("Heading 1", ""), ("Normal", "Tresc")])
    )
    assert chapters[0][1] == "Rozdzial"


# ---------------------------------------------------------------------------
# Front / back matter classification
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("title", ["Wstęp", "Przedmowa", "Introduction", "Prolog"])
def test_frontmatter_titles_are_classified(title):
    chapters, _ = parse_docx(
        F.simple_doc([("Heading 1", title), ("Normal", "x")])
    )
    assert chapters[0][2] == "frontmatter"
    assert chapters[0][3] == "doc-preface"


@pytest.mark.parametrize("title", ["Zakończenie", "Epilog", "Conclusion"])
def test_backmatter_titles_are_classified(title):
    chapters, _ = parse_docx(
        F.simple_doc([("Heading 1", title), ("Normal", "x")])
    )
    assert chapters[0][2] == "backmatter"
    assert chapters[0][3] == "doc-conclusion"


def test_classification_is_case_insensitive():
    chapters, _ = parse_docx(
        F.simple_doc([("Heading 1", "WSTĘP"), ("Normal", "x")])
    )
    assert chapters[0][2] == "frontmatter"


def test_ordinary_title_is_a_chapter():
    chapters, _ = parse_docx(
        F.simple_doc([("Heading 1", "Rozdzial pierwszy"), ("Normal", "x")])
    )
    assert chapters[0][2] == "chapter"
    assert chapters[0][3] == "doc-chapter"


# ---------------------------------------------------------------------------
# Block elements
# ---------------------------------------------------------------------------


def test_subheadings_become_h2_and_h3():
    body = parse_docx(
        F.simple_doc(
            [
                ("Heading 1", "R"),
                ("Heading 2", "Podrozdzial"),
                ("Heading 3", "Pod-pod"),
            ]
        )
    )[0][0][4]
    assert "<h2>Podrozdzial</h2>" in body
    assert "<h3>Pod-pod</h3>" in body


def test_quote_becomes_blockquote():
    body = parse_docx(
        F.simple_doc([("Heading 1", "R"), ("Quote", "Cytowane zdanie")])
    )[0][0][4]
    assert "<blockquote><p>Cytowane zdanie</p></blockquote>" in body


def test_consecutive_bullets_share_one_list():
    body = parse_docx(
        F.simple_doc(
            [
                ("Heading 1", "R"),
                ("List Bullet", "raz"),
                ("List Bullet", "dwa"),
            ]
        )
    )[0][0][4]
    assert body.count("<ul>") == 1
    assert body.count("</ul>") == 1
    assert body.count("<li>") == 2


def test_switching_from_bullets_to_numbers_closes_the_first_list():
    body = parse_docx(
        F.simple_doc(
            [
                ("Heading 1", "R"),
                ("List Bullet", "punkt"),
                ("List Number", "pozycja"),
            ]
        )
    )[0][0][4]
    assert body.index("</ul>") < body.index("<ol>")
    assert body.count("<ul>") == 1
    assert body.count("<ol>") == 1


def test_paragraph_after_list_closes_the_list():
    body = parse_docx(
        F.simple_doc(
            [
                ("Heading 1", "R"),
                ("List Bullet", "punkt"),
                ("Normal", "Zwykly akapit"),
            ]
        )
    )[0][0][4]
    assert body.index("</ul>") < body.index("<p>Zwykly akapit</p>")


def test_open_list_is_closed_at_end_of_chapter():
    body = parse_docx(
        F.simple_doc([("Heading 1", "R"), ("List Bullet", "ostatni")])
    )[0][0][4]
    assert body.rstrip().endswith("</ul>")


def test_list_is_closed_before_next_chapter_starts():
    chapters, _ = parse_docx(
        F.simple_doc(
            [
                ("Heading 1", "Pierwszy"),
                ("List Bullet", "punkt"),
                ("Heading 1", "Drugi"),
            ]
        )
    )
    assert chapters[0][4].count("<ul>") == chapters[0][4].count("</ul>") == 1
    assert "<ul>" not in chapters[1][4]


def test_truly_empty_paragraphs_produce_no_markup():
    body = parse_docx(
        F.simple_doc([("Heading 1", "R"), ("Normal", ""), ("Normal", "Tresc")])
    )[0][0][4]
    assert body == "<p>Tresc</p>"


def test_whitespace_only_paragraph_is_kept_as_a_blank_paragraph():
    """Word authors use spacer paragraphs for layout; a run holding only spaces
    is preserved rather than dropped. Valid XHTML, so EpubCheck is unaffected."""
    body = parse_docx(
        F.simple_doc([("Heading 1", "R"), ("Normal", "   "), ("Normal", "Tresc")])
    )[0][0][4]
    assert body == "<p>   </p>\n<p>Tresc</p>"


# ---------------------------------------------------------------------------
# Inline formatting
# ---------------------------------------------------------------------------


def _body_from_runs(*runs):
    doc = F.new_document()
    doc.add_paragraph("R", style="Heading 1")
    para = doc.add_paragraph()
    for text, fmt in runs:
        F.add_run(para, text, **fmt)
    return parse_docx(F.to_bytes(doc))[0][0][4]


def test_bold_and_italic_map_to_strong_and_em():
    assert "<strong>gruby</strong>" in _body_from_runs(("gruby", {"bold": True}))
    assert "<em>skosny</em>" in _body_from_runs(("skosny", {"italic": True}))


def test_bold_italic_combination_nests_em_inside_strong():
    body = _body_from_runs(("oba", {"bold": True, "italic": True}))
    assert "<strong><em>oba</em></strong>" in body


def test_underline_and_strikethrough():
    assert "<u>pod</u>" in _body_from_runs(("pod", {"underline": True}))
    assert "<s>prze</s>" in _body_from_runs(("prze", {"strike": True}))


def test_underline_wraps_bold():
    body = _body_from_runs(("mix", {"bold": True, "underline": True}))
    assert "<u><strong>mix</strong></u>" in body


def test_runs_are_concatenated_in_order():
    body = _body_from_runs(("Ala ", {}), ("ma ", {"bold": True}), ("kota", {}))
    assert "<p>Ala <strong>ma </strong>kota</p>" in body


# ---------------------------------------------------------------------------
# Escaping
# ---------------------------------------------------------------------------


def test_special_characters_are_escaped_in_body():
    body = parse_docx(
        F.simple_doc([("Heading 1", "R"), ("Normal", "Kowalski & <Syn>")])
    )[0][0][4]
    assert "&amp;" in body and "&lt;Syn&gt;" in body
    assert "<Syn>" not in body


def test_special_characters_are_escaped_in_title():
    chapters, _ = parse_docx(
        F.simple_doc([("Heading 1", "R&D <tag>"), ("Normal", "x")])
    )
    assert chapters[0][1] == "R&amp;D &lt;tag&gt;"


def test_escaped_title_still_matches_frontmatter_keyword():
    """flush_chapter unescapes before comparing, so entities must not break it."""
    chapters, _ = parse_docx(
        F.simple_doc([("Heading 1", "Wstęp"), ("Normal", "x")])
    )
    assert chapters[0][2] == "frontmatter"


# ---------------------------------------------------------------------------
# Tables
# ---------------------------------------------------------------------------


def test_table_first_row_is_header():
    doc = F.new_document()
    doc.add_paragraph("R", style="Heading 1")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Naglowek A"
    table.cell(0, 1).text = "Naglowek B"
    table.cell(1, 0).text = "Dane A"
    table.cell(1, 1).text = "Dane B"
    body = parse_docx(F.to_bytes(doc))[0][0][4]
    assert "<th>Naglowek A</th>" in body
    assert "<td>Dane A</td>" in body
    assert "<th>Dane A</th>" not in body


def test_table_cell_content_is_escaped():
    doc = F.new_document()
    doc.add_paragraph("R", style="Heading 1")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "a & b"
    assert "a &amp; b" in parse_docx(F.to_bytes(doc))[0][0][4]


def test_table_before_any_heading_is_ignored():
    doc = F.new_document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "sierota"
    doc.add_paragraph("R", style="Heading 1")
    chapters, _ = parse_docx(F.to_bytes(doc))
    assert "sierota" not in chapters[0][4]


def test_table_to_html_is_balanced():
    doc = F.new_document()
    table = doc.add_table(rows=3, cols=2)
    html = table_to_html(table)
    assert html.count("<tr>") == html.count("</tr>") == 3
    assert html.startswith("<table>") and html.endswith("</table>")


# ---------------------------------------------------------------------------
# Hyperlinks
# ---------------------------------------------------------------------------


def test_hyperlink_becomes_anchor_with_href():
    doc = F.new_document()
    doc.add_paragraph("R", style="Heading 1")
    F.add_hyperlink(doc.add_paragraph(), "https://example.com/a?x=1&y=2", "Klik")
    body = parse_docx(F.to_bytes(doc))[0][0][4]
    assert 'href="https://example.com/a?x=1&amp;y=2"' in body
    assert ">Klik</a>" in body


def test_hyperlink_text_survives_when_relationship_is_missing():
    doc = F.new_document()
    doc.add_paragraph("R", style="Heading 1")
    para = doc.add_paragraph()
    F.add_hyperlink(para, "https://example.com", "Widoczny tekst")
    # Drop the r:id so the target cannot be resolved.
    from docx.oxml.ns import qn

    para._p.findall(qn("w:hyperlink"))[0].attrib.pop(qn("r:id"))
    body = parse_docx(F.to_bytes(doc))[0][0][4]
    assert "Widoczny tekst" in body
    assert "<a href=" not in body


# ---------------------------------------------------------------------------
# Images
# ---------------------------------------------------------------------------


def test_inline_image_is_extracted_and_referenced():
    doc = F.new_document()
    doc.add_paragraph("R", style="Heading 1")
    F.add_image(doc.add_paragraph())
    chapters, images = parse_docx(F.to_bytes(doc))
    assert list(images) == ["img001.png"]
    assert images["img001.png"].startswith(b"\x89PNG")
    assert 'src="../images/img001.png"' in chapters[0][4]


def test_images_are_numbered_across_chapters():
    doc = F.new_document()
    doc.add_paragraph("R1", style="Heading 1")
    F.add_image(doc.add_paragraph())
    doc.add_paragraph("R2", style="Heading 1")
    F.add_image(doc.add_paragraph())
    _, images = parse_docx(F.to_bytes(doc))
    assert sorted(images) == ["img001.png", "img002.png"]


def test_image_uses_docpr_description_as_alt_text():
    doc = F.new_document()
    doc.add_paragraph("R", style="Heading 1")
    F.add_image(doc.add_paragraph(), alt="Wykres sprzedazy")
    assert 'alt="Wykres sprzedazy"' in parse_docx(F.to_bytes(doc))[0][0][4]


def test_image_alt_text_is_escaped():
    doc = F.new_document()
    doc.add_paragraph("R", style="Heading 1")
    F.add_image(doc.add_paragraph(), alt='Rys. "A" & B')
    body = parse_docx(F.to_bytes(doc))[0][0][4]
    assert "&amp;" in body
    assert 'alt="Rys. "A" & B"' not in body


@pytest.mark.parametrize(
    "lang,expected", [("pl", "Ilustracja"), ("en", "Illustration"), ("de", "Abbildung")]
)
def test_image_without_description_uses_localized_default_alt(lang, expected):
    doc = F.new_document()
    doc.add_paragraph("R", style="Heading 1")
    F.add_image(doc.add_paragraph())
    assert f'alt="{expected}"' in parse_docx(F.to_bytes(doc), lang)[0][0][4]


def test_unknown_language_falls_back_to_english_alt():
    doc = F.new_document()
    doc.add_paragraph("R", style="Heading 1")
    F.add_image(doc.add_paragraph())
    assert 'alt="Illustration"' in parse_docx(F.to_bytes(doc), "xx")[0][0][4]


# ---------------------------------------------------------------------------
# Footnotes
# ---------------------------------------------------------------------------


def test_extract_footnotes_reads_a_generic_part():
    """python-docx loads footnotes.xml as a plain Part with no parsed element,
    so extraction must fall back to parsing the raw blob."""
    doc = F.new_document()
    F.attach_footnotes(doc, {"2": "Tresc przypisu"})
    doc.add_paragraph("R", style="Heading 1")
    reloaded = Document(io.BytesIO(F.to_bytes(doc)))
    assert extract_footnotes(reloaded) == {"2": "Tresc przypisu"}


def test_separator_footnotes_are_skipped():
    doc = F.new_document()
    F.attach_footnotes(doc, {"3": "Prawdziwy przypis"})
    doc.add_paragraph("R", style="Heading 1")
    notes = extract_footnotes(Document(io.BytesIO(F.to_bytes(doc))))
    assert set(notes) == {"3"}


def test_document_without_footnotes_part_returns_empty_mapping():
    doc = F.new_document()
    doc.add_paragraph("R", style="Heading 1")
    assert extract_footnotes(Document(io.BytesIO(F.to_bytes(doc)))) == {}


def test_footnote_reference_emits_noteref_and_matching_aside():
    doc = F.new_document()
    F.attach_footnotes(doc, {"2": "Zrodlo: GUS"})
    doc.add_paragraph("R", style="Heading 1")
    F.add_footnote_reference(doc.add_paragraph("Zdanie"), "2")
    body = parse_docx(F.to_bytes(doc))[0][0][4]
    assert '<sup><a epub:type="noteref" href="#fn-2">[2]</a></sup>' in body
    assert 'id="fn-2"' in body
    assert "Zrodlo: GUS" in body
    assert 'role="note"' in body


def test_footnote_section_is_emitted_once_per_chapter():
    doc = F.new_document()
    F.attach_footnotes(doc, {"2": "Nota"})
    doc.add_paragraph("R", style="Heading 1")
    F.add_footnote_reference(doc.add_paragraph("A"), "2")
    F.add_footnote_reference(doc.add_paragraph("B"), "2")
    body = parse_docx(F.to_bytes(doc))[0][0][4]
    assert body.count('class="footnotes-section"') == 1
    assert body.count('id="fn-2"') == 1
    assert body.count('href="#fn-2"') == 2


def test_footnotes_do_not_leak_between_chapters():
    doc = F.new_document()
    F.attach_footnotes(doc, {"2": "Nota pierwsza", "3": "Nota druga"})
    doc.add_paragraph("R1", style="Heading 1")
    F.add_footnote_reference(doc.add_paragraph("A"), "2")
    doc.add_paragraph("R2", style="Heading 1")
    F.add_footnote_reference(doc.add_paragraph("B"), "3")
    chapters, _ = parse_docx(F.to_bytes(doc))
    assert "Nota pierwsza" in chapters[0][4]
    assert "Nota druga" not in chapters[0][4]
    assert "Nota druga" in chapters[1][4]
    assert "Nota pierwsza" not in chapters[1][4]


def test_chapter_without_footnotes_has_no_footnote_section():
    doc = F.new_document()
    F.attach_footnotes(doc, {"2": "Nota"})
    doc.add_paragraph("R", style="Heading 1")
    doc.add_paragraph("Bez przypisow")
    assert "footnotes-section" not in parse_docx(F.to_bytes(doc))[0][0][4]


def test_reference_to_unknown_footnote_does_not_emit_an_aside():
    doc = F.new_document()
    F.attach_footnotes(doc, {"2": "Nota"})
    doc.add_paragraph("R", style="Heading 1")
    F.add_footnote_reference(doc.add_paragraph("Zdanie"), "99")
    body = parse_docx(F.to_bytes(doc))[0][0][4]
    assert 'href="#fn-99"' in body
    assert "footnotes-section" not in body


def test_every_noteref_target_exists_in_the_same_chapter():
    """Guards against dangling internal links, which EpubCheck rejects."""
    import re

    doc = F.new_document()
    F.attach_footnotes(doc, {"2": "Nota A", "3": "Nota B"})
    doc.add_paragraph("R", style="Heading 1")
    F.add_footnote_reference(doc.add_paragraph("A"), "2")
    F.add_footnote_reference(doc.add_paragraph("B"), "3")
    body = parse_docx(F.to_bytes(doc))[0][0][4]
    targets = set(re.findall(r'href="#([^"]+)"', body))
    ids = set(re.findall(r'id="([^"]+)"', body))
    assert targets and targets <= ids
